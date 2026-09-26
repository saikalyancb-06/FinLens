"""Ingestion tests for the B2B API: detection, parsers, and safe intake.

Every fixture is generated here rather than checked in, so a test states the
exact bytes it is making a claim about, and every assertion is on a *parsed
value* — paise, direction, date, statement metadata — not on "it did not
raise". A parser that silently returns zero transactions passes a
does-not-raise test and fails every one of these.

The two currencies in play are deliberate: the CSV/TSV/JSON/XLSX fixtures are
Indian statements in rupees, the OFX and CAMT fixtures are what those formats
actually arrive as, so the multi-currency and per-row-currency paths are
exercised on real shapes rather than on an INR file with a relabelled column.
"""
from __future__ import annotations

import datetime
import io
import json
import os
import zipfile

import pytest

from app.b2b.canonical import CREDIT, DEBIT
from app.b2b.detect import detect_format
from app.b2b.errors import ApiError
from app.b2b.ingest import IngestedFile, cleanup, sanitise_filename, save_upload
from app.b2b.metrics import (
    W_CONTINUITY_UNVERIFIABLE,
    W_MULTI_CURRENCY,
)
from app.b2b.parsers import delimited, legacy, ofx
from app.b2b.parsers.registry import (
    SUPPORTED_FORMATS,
    UNSUPPORTED_FORMATS,
    get_parser,
)

# ---------------------------------------------------------------------------
# Fixture content. One statement, expressed in every format under test, so a
# cross-format disagreement shows up as a failing equality rather than as two
# separately-passing tests.
# ---------------------------------------------------------------------------

ROWS = [
    # date,        description,          debit,      credit,     balance
    ("01/04/2024", "UPI-SWIGGY-123456", "500.50", "", "10000.00"),
    ("02/04/2024", "SALARY CREDIT ACME", "", "25000.00", "35000.00"),
    ("03/04/2024", "NEFT-RENT PAYMENT", "12000.00", "", "23000.00"),
]

#: What every tabular fixture must produce: (date, debit_paise, credit_paise,
#: balance_paise, direction).
EXPECTED = [
    (datetime.date(2024, 4, 1), 50050, None, 1000000, DEBIT),
    (datetime.date(2024, 4, 2), None, 2500000, 3500000, CREDIT),
    (datetime.date(2024, 4, 3), 1200000, None, 2300000, DEBIT),
]

HEADER = ["Date", "Description", "Debit", "Credit", "Balance"]


def _write(path, text: str, encoding: str = "utf-8") -> str:
    with open(path, "w", encoding=encoding, newline="") as handle:
        handle.write(text)
    return str(path)


def make_delimited(path, delimiter: str) -> str:
    lines = [delimiter.join(HEADER)]
    lines += [delimiter.join(row) for row in ROWS]
    return _write(path, "\n".join(lines) + "\n")


def make_json(path, shape: str) -> str:
    records = [
        {"date": "2024-04-01", "description": "UPI-SWIGGY-123456",
         "debit": 500.50, "credit": 0, "balance": 10000.00},
        {"date": "2024-04-02", "description": "SALARY CREDIT ACME",
         "debit": 0, "credit": 25000.00, "balance": 35000.00},
        {"date": "2024-04-03", "description": "NEFT-RENT PAYMENT",
         "debit": 12000.00, "credit": 0, "balance": 23000.00},
    ]
    if shape == "list":
        document = records
    elif shape == "wrapper":
        document = {"transactions": records}
    elif shape == "nested":
        document = {"data": {"transactions": records}, "status": "ok"}
    else:  # pragma: no cover - guards a typo in a test
        raise AssertionError(f"unknown shape {shape}")
    return _write(path, json.dumps(document, indent=1))


OFX1_SGML = """OFXHEADER:100
DATA:OFXSGML
VERSION:102
SECURITY:NONE
ENCODING:USASCII
CHARSET:1252
COMPRESSION:NONE
OLDFILEUID:NONE
NEWFILEUID:NONE

<OFX>
<BANKMSGSRSV1>
<STMTTRNRS>
<TRNUID>1001
<STMTRS>
<CURDEF>INR
<BANKACCTFROM>
<BANKID>HDFC0000123
<ACCTID>000123456789
<ACCTTYPE>SAVINGS
</BANKACCTFROM>
<BANKTRANLIST>
<DTSTART>20240401000000.000[+5:IST]
<DTEND>20240430235959.000[+5:IST]
<STMTTRN>
<TRNTYPE>DEBIT
<DTPOSTED>20240401120000.000[+5:IST]
<TRNAMT>-1250.75
<FITID>FIT0001
<NAME>SWIGGY BANGALORE
<MEMO>UPI PAYMENT
</STMTTRN>
<STMTTRN>
<TRNTYPE>DIRECTDEP
<DTPOSTED>20240405
<TRNAMT>78000.00
<FITID>FIT0002
<NAME>ACME PAYROLL
</STMTTRN>
</BANKTRANLIST>
<LEDGERBAL>
<BALAMT>96749.25
<DTASOF>20240430
</LEDGERBAL>
</STMTRS>
</STMTTRNRS>
</BANKMSGSRSV1>
</OFX>
"""

OFX2_XML = """<?xml version="1.0" encoding="UTF-8"?>
<?OFX OFXHEADER="200" VERSION="211" SECURITY="NONE" OLDFILEUID="NONE" NEWFILEUID="NONE"?>
<OFX>
  <BANKMSGSRSV1>
    <STMTTRNRS>
      <TRNUID>2001</TRNUID>
      <STMTRS>
        <CURDEF>USD</CURDEF>
        <BANKACCTFROM>
          <BANKID>021000021</BANKID>
          <ACCTID>987654321000</ACCTID>
          <ACCTTYPE>CHECKING</ACCTTYPE>
        </BANKACCTFROM>
        <BANKTRANLIST>
          <DTSTART>20240301000000</DTSTART>
          <DTEND>20240331000000</DTEND>
          <STMTTRN>
            <TRNTYPE>POS</TRNTYPE>
            <DTPOSTED>20240302083000.000[-5:EST]</DTPOSTED>
            <TRNAMT>-42.10</TRNAMT>
            <FITID>XFIT-1</FITID>
            <NAME>WHOLE FOODS</NAME>
            <MEMO>CARD 4411</MEMO>
          </STMTTRN>
          <STMTTRN>
            <TRNTYPE>XFER</TRNTYPE>
            <DTPOSTED>20240315000000</DTPOSTED>
            <TRNAMT>1500.00</TRNAMT>
            <FITID>XFIT-2</FITID>
            <NAME>TRANSFER FROM SAVINGS</NAME>
          </STMTTRN>
        </BANKTRANLIST>
        <LEDGERBAL>
          <BALAMT>3457.90</BALAMT>
          <DTASOF>20240331</DTASOF>
        </LEDGERBAL>
      </STMTRS>
    </STMTTRNRS>
  </BANKMSGSRSV1>
</OFX>
"""

CAMT053_XML = """<?xml version="1.0" encoding="UTF-8"?>
<Document xmlns="urn:iso:std:iso:20022:tech:xsd:camt.053.001.08">
  <BkToCstmrStmt>
    <GrpHdr>
      <MsgId>MSG-2024-04</MsgId>
      <CreDtTm>2024-05-01T06:00:00</CreDtTm>
    </GrpHdr>
    <Stmt>
      <Id>STMT-0001</Id>
      <Acct>
        <Id><IBAN>DE89370400440532013000</IBAN></Id>
        <Ccy>EUR</Ccy>
      </Acct>
      <FrToDt>
        <FrDtTm>2024-04-01T00:00:00</FrDtTm>
        <ToDtTm>2024-04-30T23:59:59</ToDtTm>
      </FrToDt>
      <Bal>
        <Tp><CdOrPrtry><Cd>OPBD</Cd></CdOrPrtry></Tp>
        <Amt Ccy="EUR">1000.00</Amt>
        <CdtDbtInd>CRDT</CdtDbtInd>
        <Dt><Dt>2024-04-01</Dt></Dt>
      </Bal>
      <Bal>
        <Tp><CdOrPrtry><Cd>CLBD</Cd></CdOrPrtry></Tp>
        <Amt Ccy="EUR">1750.25</Amt>
        <CdtDbtInd>CRDT</CdtDbtInd>
        <Dt><Dt>2024-04-30</Dt></Dt>
      </Bal>
      <Ntry>
        <NtryRef>NTRY-1</NtryRef>
        <Amt Ccy="EUR">250.75</Amt>
        <CdtDbtInd>DBIT</CdtDbtInd>
        <Sts><Cd>BOOK</Cd></Sts>
        <BookgDt><Dt>2024-04-05</Dt></BookgDt>
        <ValDt><Dt>2024-04-06</Dt></ValDt>
        <BkTxCd><Prtry><Cd>PMNT-CCRD-POSD</Cd></Prtry></BkTxCd>
        <AddtlNtryInf>CARD PURCHASE MEDIAMARKT BERLIN</AddtlNtryInf>
      </Ntry>
      <Ntry>
        <AcctSvcrRef>ASR-0002</AcctSvcrRef>
        <Amt Ccy="EUR">1001.00</Amt>
        <CdtDbtInd>CRDT</CdtDbtInd>
        <BookgDt><Dt>2024-04-10</Dt></BookgDt>
        <NtryDtls>
          <TxDtls>
            <RmtInf><Ustrd>INVOICE 2024-114</Ustrd></RmtInf>
          </TxDtls>
        </NtryDtls>
      </Ntry>
    </Stmt>
  </BkToCstmrStmt>
</Document>
"""

#: Same schema, two currencies, to exercise the per-row currency path.
CAMT053_MULTICCY = """<?xml version="1.0" encoding="UTF-8"?>
<Document xmlns="urn:iso:std:iso:20022:tech:xsd:camt.053.001.02">
  <BkToCstmrStmt>
    <Stmt>
      <Id>STMT-FX</Id>
      <Acct><Id><Othr><Id>NL02ABNA0123456789</Id></Othr></Id></Acct>
      <Ntry>
        <Amt Ccy="EUR">100.00</Amt>
        <CdtDbtInd>DBIT</CdtDbtInd>
        <BookgDt><Dt>2024-04-02</Dt></BookgDt>
        <AddtlNtryInf>EUR FEE</AddtlNtryInf>
      </Ntry>
      <Ntry>
        <Amt Ccy="USD">250.00</Amt>
        <CdtDbtInd>CRDT</CdtDbtInd>
        <BookgDt><Dt>2024-04-03</Dt></BookgDt>
        <AddtlNtryInf>USD RECEIPT</AddtlNtryInf>
      </Ntry>
    </Stmt>
  </BkToCstmrStmt>
</Document>
"""

#: An Ntry whose details are individually priced. Documented behaviour is one
#: transaction per TxDtls in exactly this case, and one per Ntry otherwise.
CAMT053_BATCH = """<?xml version="1.0" encoding="UTF-8"?>
<Document xmlns="urn:iso:std:iso:20022:tech:xsd:camt.053.001.08">
  <BkToCstmrStmt>
    <Stmt>
      <Id>STMT-BATCH</Id>
      <Ntry>
        <NtryRef>BATCH-1</NtryRef>
        <Amt Ccy="EUR">300.00</Amt>
        <CdtDbtInd>DBIT</CdtDbtInd>
        <BookgDt><Dt>2024-04-08</Dt></BookgDt>
        <NtryDtls>
          <TxDtls>
            <Amt Ccy="EUR">120.00</Amt>
            <CdtDbtInd>DBIT</CdtDbtInd>
            <RmtInf><Ustrd>SUPPLIER A</Ustrd></RmtInf>
          </TxDtls>
          <TxDtls>
            <Amt Ccy="EUR">180.00</Amt>
            <CdtDbtInd>DBIT</CdtDbtInd>
            <RmtInf><Ustrd>SUPPLIER B</Ustrd></RmtInf>
          </TxDtls>
        </NtryDtls>
      </Ntry>
    </Stmt>
  </BkToCstmrStmt>
</Document>
"""


def make_xlsx(path) -> str:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Statement"
    sheet.append(HEADER)
    for row in ROWS:
        sheet.append(list(row))
    workbook.save(path)
    return str(path)


def make_pdf(path) -> str:
    """A one-page PDF of the same statement, drawn as a bordered table."""
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=10)
    pdf.cell(0, 8, "ACME BANK - ACCOUNT STATEMENT", ln=1)
    widths = (28, 60, 26, 28, 30)
    for row in [tuple(HEADER)] + [tuple(r) for r in ROWS]:
        for width, value in zip(widths, row):
            pdf.cell(width, 7, value, border=1)
        pdf.ln()
    pdf.output(str(path))
    return str(path)


class FakeUpload:
    """The two attributes `save_upload` touches on a FastAPI UploadFile."""

    def __init__(self, filename: str, data: bytes):
        self.filename = filename
        self.file = io.BytesIO(data)


def assert_matches_expected(transactions, *, currency="INR"):
    """Every tabular fixture must land on exactly these values."""
    assert len(transactions) == len(EXPECTED)
    for txn, (date, debit, credit, balance, direction) in zip(transactions, EXPECTED):
        assert txn.txn_date == date
        assert txn.debit_paise == debit
        assert txn.credit_paise == credit
        assert txn.balance_paise == balance
        assert txn.direction == direction
        assert txn.currency == currency
    assert transactions[0].narration_raw == "UPI-SWIGGY-123456"
    assert transactions[1].narration_raw == "SALARY CREDIT ACME"
    # Money is integers all the way through; a float here means a rounding bug.
    assert all(isinstance(t.amount_paise, int) for t in transactions)


# ===========================================================================
# detect.py
# ===========================================================================

def test_detect_pdf_by_magic_bytes(tmp_path):
    path = tmp_path / "statement.pdf"
    path.write_bytes(b"%PDF-1.7\n%\xe2\xe3\xcf\xd3\n1 0 obj\n")
    detected = detect_format("statement.pdf", path.read_bytes(), str(path))
    assert detected.format == "pdf"
    assert detected.confidence == 1.0
    assert detected.mismatch is False
    assert detected.mime == "application/pdf"


def test_detect_extension_content_mismatch_is_reported(tmp_path):
    """A PDF named .csv is a PDF, and the API is told the name lied."""
    path = tmp_path / "statement.csv"
    path.write_bytes(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\ntrailer\n")
    detected = detect_format("statement.csv", path.read_bytes(), str(path))
    assert detected.format == "pdf", "content must win over the extension"
    assert detected.mismatch is True
    assert detected.extension_format == "csv"
    assert "extension '.csv' claimed csv" in detected.reason


def test_detect_xlsx_probes_zip_members(tmp_path):
    path = make_xlsx(tmp_path / "book.xlsx")
    with open(path, "rb") as handle:
        head = handle.read(8192)
    detected = detect_format("book.xlsx", head, path)
    assert detected.format == "xlsx"
    assert detected.confidence == 1.0
    assert "xl/" in detected.reason


def test_detect_plain_zip_is_not_xlsx(tmp_path):
    path = tmp_path / "statements.zip"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("april.csv", "Date,Description\n")
    head = path.read_bytes()[:8192]
    detected = detect_format("statements.zip", head, str(path))
    assert detected.format == "zip"
    assert detected.is_supported is False


def test_detect_zip_renamed_to_xlsx_is_still_zip(tmp_path):
    path = tmp_path / "book.xlsx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("readme.txt", "not a workbook")
    detected = detect_format("book.xlsx", path.read_bytes()[:8192], str(path))
    assert detected.format == "zip"
    assert detected.mismatch is True
    assert detected.extension_format == "xlsx"


def test_detect_ole2_is_xls(tmp_path):
    path = tmp_path / "old.xls"
    path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64)
    detected = detect_format("old.xls", path.read_bytes(), str(path))
    assert detected.format == "xls"


def test_detect_delimiters(tmp_path):
    comma = make_delimited(tmp_path / "a.csv", ",")
    semi = make_delimited(tmp_path / "b.csv", ";")
    tab = make_delimited(tmp_path / "c.tsv", "\t")
    pipe = make_delimited(tmp_path / "d.txt", "|")

    assert detect_format("a.csv", open(comma, "rb").read(), comma).delimiter == ","
    semi_detected = detect_format("b.csv", open(semi, "rb").read(), semi)
    assert semi_detected.format == "csv"
    assert semi_detected.delimiter == ";"
    # A delimiter difference is not a lie about the format.
    assert semi_detected.mismatch is False
    assert detect_format("c.tsv", open(tab, "rb").read(), tab).format == "tsv"
    assert detect_format("d.txt", open(pipe, "rb").read(), pipe).delimiter == "|"


def test_detect_ofx_versions_and_qfx(tmp_path):
    ofx1 = _write(tmp_path / "s.ofx", OFX1_SGML)
    ofx2 = _write(tmp_path / "s2.ofx", OFX2_XML)
    qfx = _write(tmp_path / "s.qfx", OFX1_SGML)

    first = detect_format("s.ofx", open(ofx1, "rb").read(), ofx1)
    assert first.format == "ofx"
    assert first.detail["ofx_version"] == "1.x"

    second = detect_format("s2.ofx", open(ofx2, "rb").read(), ofx2)
    assert second.format == "ofx"
    assert second.detail["ofx_version"] == "2.x", "OFX 2.x must not be read as CAMT"

    # Only the filename separates QFX from OFX; both parse identically.
    assert detect_format("s.qfx", open(qfx, "rb").read(), qfx).format == "qfx"


def test_detect_camt_and_reject_other_xml(tmp_path):
    camt_path = _write(tmp_path / "camt.xml", CAMT053_XML)
    detected = detect_format("camt.xml", open(camt_path, "rb").read(), camt_path)
    assert detected.format == "xml_camt"

    other = _write(tmp_path / "other.xml", "<?xml version='1.0'?>\n<invoices><i/></invoices>")
    other_detected = detect_format("other.xml", open(other, "rb").read(), other)
    assert other_detected.format == "unknown"
    assert other_detected.detail["xml_root"] == "invoices"
    assert "CAMT.053" in other_detected.reason


def test_detect_json_and_broken_json(tmp_path):
    good = make_json(tmp_path / "s.json", "wrapper")
    assert detect_format("s.json", open(good, "rb").read(), good).format == "json"

    broken = _write(tmp_path / "broken.json", '{"transactions": [ {"date": ')
    detected = detect_format("broken.json", open(broken, "rb").read(), broken)
    assert detected.format == "json"
    assert detected.confidence < 0.6, "an unparseable JSON must not claim confidence"


def test_detect_empty_input_is_unknown():
    detected = detect_format("nothing.csv", b"")
    assert detected.format == "unknown"
    assert detected.confidence == 0.0


# ===========================================================================
# Tabular parsers: CSV, TSV, semicolon CSV, XLSX
# ===========================================================================

def test_csv_parses_exact_values(tmp_path):
    path = make_delimited(tmp_path / "statement.csv", ",")
    output = get_parser("csv")(path)

    assert_matches_expected(output.transactions)
    assert output.transactions[0].reference_number == "SWIGGY"
    assert output.rows_checked_for_continuity == 2
    assert output.continuity_pass_rate == 1.0
    assert output.continuity_passed is True
    assert W_CONTINUITY_UNVERIFIABLE not in output.warning_codes

    meta = output.statement_meta
    assert meta["closing_balance"] == 23000.00
    assert meta["closing_balance_paise"] == 2300000
    assert meta["closing_balance_basis"] == "last_row_balance"
    assert meta["period_start"] == "2024-04-01"
    assert meta["period_end"] == "2024-04-03"
    assert meta["period_source"] == "observed"
    assert meta["currency"] == "INR"


def test_tsv_parses_exact_values(tmp_path):
    path = make_delimited(tmp_path / "statement.tsv", "\t")
    output = get_parser("tsv")(path)

    assert_matches_expected(output.transactions)
    assert output.statement_meta["delimiter"] == "\t"
    assert output.rows_checked_for_continuity == 2
    assert output.continuity_passed is True


def test_pipe_delimited_txt_parses_exact_values(tmp_path):
    path = make_delimited(tmp_path / "statement.txt", "|")
    output = get_parser("txt")(path)

    assert_matches_expected(output.transactions)
    assert output.statement_meta["delimiter"] == "|"


def test_semicolon_matches_comma_equivalent(tmp_path):
    """The point of sniffing: the same statement, two separators, one result.

    The comma file goes through the legacy pipeline and the semicolon file
    through the delimited parser, so this also asserts the two code paths agree
    rather than merely that each is self-consistent.
    """
    comma = get_parser("csv")(make_delimited(tmp_path / "comma.csv", ","))
    semi = get_parser("csv")(make_delimited(tmp_path / "semi.csv", ";"))

    assert semi.statement_meta["delimiter"] == ";"
    assert len(semi.transactions) == len(comma.transactions) == 3

    def shape(txn):
        return (txn.txn_date, txn.direction, txn.debit_paise, txn.credit_paise,
                txn.balance_paise, txn.narration_raw, txn.reference_number,
                txn.row_index)

    assert [shape(t) for t in semi.transactions] == [shape(t) for t in comma.transactions]
    assert semi.rows_checked_for_continuity == comma.rows_checked_for_continuity == 2


def test_csv_dispatch_routes_comma_to_legacy_and_the_rest_to_delimited(tmp_path, monkeypatch):
    """The comma path keeps its years of production handling; the rest do not need it."""
    seen = []
    monkeypatch.setattr(legacy, "parse",
                        lambda path, **kwargs: seen.append(("legacy", path)))
    monkeypatch.setattr(delimited, "parse",
                        lambda path, **kwargs: seen.append(("delimited", path)))

    get_parser("csv")(make_delimited(tmp_path / "comma.csv", ","))
    get_parser("csv")(make_delimited(tmp_path / "semi.csv", ";"))
    get_parser("csv")(make_delimited(tmp_path / "tabbed.csv", "\t"))

    assert [kind for kind, _ in seen] == ["legacy", "delimited", "delimited"]


def test_semicolon_csv_through_legacy_path_would_have_found_nothing(tmp_path):
    """Why the dispatch exists: the comma reader sees one column, not five."""
    path = make_delimited(tmp_path / "semi.csv", ";")
    grid = [line.split(",") for line in open(path).read().splitlines()]
    header_idx, mapping = delimited.find_header(grid)
    assert mapping == {}, "a comma split of a semicolon CSV maps no headers"
    assert header_idx == -1


def test_xlsx_parses_exact_values(tmp_path):
    path = make_xlsx(tmp_path / "statement.xlsx")
    output = get_parser("xlsx")(path)

    assert_matches_expected(output.transactions)
    assert output.rows_checked_for_continuity == 2
    assert output.continuity_passed is True
    assert output.statement_meta["closing_balance"] == 23000.00


def test_pdf_parses_exact_values(tmp_path):
    """The legacy PDF path, exercised through the adapter rather than trusted."""
    pytest.importorskip("fpdf", reason="fpdf2 builds the PDF fixture")
    path = make_pdf(tmp_path / "statement.pdf")

    detected = detect_format("statement.pdf", open(path, "rb").read(8192), path)
    assert detected.format == "pdf"

    output = get_parser(detected.format)(path)
    assert_matches_expected(output.transactions)
    assert output.rows_checked_for_continuity == 2
    assert output.continuity_passed is True
    assert output.statement_meta["closing_balance"] == 23000.00
    assert output.transactions[0].source_format == "pdf"


def test_continuity_is_unverifiable_without_a_balance_column(tmp_path):
    """The distinction the legacy pipeline hides: no checks is not all-passed."""
    text = "Date,Description,Debit,Credit\n" + "\n".join(
        ",".join(row[:4]) for row in ROWS) + "\n"
    path = _write(tmp_path / "nobalance.csv", text)

    output = get_parser("csv")(path)

    assert len(output.transactions) == 3
    assert output.rows_checked_for_continuity == 0
    assert output.continuity_pass_rate is None, (
        "the pipeline reports 1.0 here; a null is the only honest answer")
    assert output.continuity_passed is None
    assert W_CONTINUITY_UNVERIFIABLE in output.warning_codes
    assert all(t.balance_paise is None for t in output.transactions)


def test_broken_balance_column_is_reported_not_hidden(tmp_path):
    text = (
        "Date,Description,Debit,Credit,Balance\n"
        "01/04/2024,OPENING SPEND,500.00,,10000.00\n"
        "02/04/2024,SECOND SPEND,500.00,,42000.00\n"
    )
    path = _write(tmp_path / "jumpy.csv", text)
    output = get_parser("csv")(path)

    assert output.rows_checked_for_continuity == 1
    assert output.continuity_passed is False
    assert output.continuity_pass_rate == 0.0
    assert "BALANCE_CONTINUITY_FAILED" in output.warning_codes
    assert output.transactions[1].balance_anomaly is True


# ===========================================================================
# JSON
# ===========================================================================

@pytest.mark.parametrize("shape", ["list", "wrapper", "nested"])
def test_json_all_three_envelopes(tmp_path, shape):
    path = make_json(tmp_path / f"{shape}.json", shape)
    output = get_parser("json")(path)

    assert_matches_expected(output.transactions)
    assert output.rows_checked_for_continuity == 2
    assert output.continuity_passed is True
    assert output.statement_meta["closing_balance"] == 23000.00


def test_json_signed_amounts_set_direction(tmp_path):
    """A feed with no type column: the sign is the direction."""
    document = {"transactions": [
        {"date": "2024-04-01", "description": "CARD SPEND", "amount": -1250.75,
         "balance": 8749.25},
        {"date": "2024-04-02", "description": "REFUND", "amount": 250.00,
         "balance": 8999.25},
    ]}
    path = _write(tmp_path / "signed.json", json.dumps(document))
    output = get_parser("json")(path)

    assert len(output.transactions) == 2
    first, second = output.transactions
    assert first.direction == DEBIT
    assert first.debit_paise == 125075
    assert first.credit_paise is None
    assert second.direction == CREDIT
    assert second.credit_paise == 25000
    assert second.balance_paise == 899925


def test_json_explicit_type_field_and_per_row_currency(tmp_path):
    document = [
        {"date": "2024-04-01", "description": "EUR FEE", "amount": 100.00,
         "type": "debit", "balance": 900.00, "currency": "EUR"},
        {"date": "2024-04-02", "description": "USD RECEIPT", "amount": 250.00,
         "type": "credit", "balance": 1150.00, "currency": "USD"},
    ]
    path = _write(tmp_path / "fx.json", json.dumps(document))
    output = get_parser("json")(path)

    assert [t.currency for t in output.transactions] == ["EUR", "USD"]
    assert output.transactions[0].direction == DEBIT
    assert output.transactions[0].debit_paise == 10000
    assert output.transactions[1].direction == CREDIT
    assert output.transactions[1].credit_paise == 25000
    assert W_MULTI_CURRENCY in output.warning_codes
    assert output.statement_meta["currencies"] == ["EUR", "USD"]


def test_json_nested_objects_are_flattened_to_headers(tmp_path):
    document = {"transactions": [
        {"txn": {"date": "2024-04-01", "narration": "UPI-SWIGGY-123456"},
         "amounts": {"debit": 500.50, "credit": 0, "balance": 10000.00},
         "tags": ["food", "upi"]},
    ]}
    path = _write(tmp_path / "nested_fields.json", json.dumps(document))
    output = get_parser("json")(path)

    assert len(output.transactions) == 1
    txn = output.transactions[0]
    assert txn.txn_date == datetime.date(2024, 4, 1)
    assert txn.debit_paise == 50050
    assert txn.narration_raw == "UPI-SWIGGY-123456"


def test_json_unknown_envelope_is_rejected_clearly(tmp_path):
    path = _write(tmp_path / "weird.json", json.dumps({"rowsOfStuff": {"a": 1}}))
    with pytest.raises(ApiError) as excinfo:
        get_parser("json")(path)
    assert excinfo.value.code == "PARSE_FAILED"
    assert "transactions" in excinfo.value.message


def test_json_corrupt_document_is_file_corrupt(tmp_path):
    path = _write(tmp_path / "broken.json", '{"transactions": [ {"date": "2024-')
    with pytest.raises(ApiError) as excinfo:
        get_parser("json")(path)
    assert excinfo.value.code == "FILE_CORRUPT"
    assert excinfo.value.status_code == 422


# ===========================================================================
# OFX / QFX
# ===========================================================================

def test_ofx_1x_sgml_parses_exact_values(tmp_path):
    path = _write(tmp_path / "statement.ofx", OFX1_SGML)
    output = get_parser("ofx")(path)

    assert len(output.transactions) == 2
    debit, credit = output.transactions

    assert debit.txn_date == datetime.date(2024, 4, 1)
    assert debit.direction == DEBIT
    assert debit.debit_paise == 125075
    assert debit.credit_paise is None
    assert debit.reference_number == "FIT0001"
    assert debit.narration_raw == "SWIGGY BANGALORE UPI PAYMENT"
    assert debit.transaction_method == "DEBIT"
    assert debit.currency == "INR"

    # TRNTYPE DIRECTDEP carries no direction of its own; the positive sign does.
    assert credit.txn_date == datetime.date(2024, 4, 5)
    assert credit.direction == CREDIT
    assert credit.credit_paise == 7800000
    assert credit.debit_paise is None
    assert credit.reference_number == "FIT0002"
    assert credit.transaction_method == "DIRECTDEP"

    meta = output.statement_meta
    assert meta["closing_balance"] == 96749.25
    assert meta["closing_balance_paise"] == 9674925
    assert meta["closing_balance_basis"] == "reported"
    assert meta["period_start"] == "2024-04-01"
    assert meta["period_end"] == "2024-04-30"
    assert meta["period_source"] == "declared"
    assert meta["account_number_masked"] == "********6789"
    assert meta["currency"] == "INR"


def test_ofx_has_no_running_balance_and_says_so(tmp_path):
    path = _write(tmp_path / "statement.ofx", OFX1_SGML)
    output = get_parser("ofx")(path)

    assert all(t.balance_paise is None for t in output.transactions)
    assert output.rows_checked_for_continuity == 0
    assert output.continuity_passed is None, "OFX must never claim a clean continuity"
    assert output.continuity_pass_rate is None
    assert W_CONTINUITY_UNVERIFIABLE in output.warning_codes


def test_ofx_2x_xml_parses_exact_values(tmp_path):
    path = _write(tmp_path / "statement.ofx", OFX2_XML)
    output = get_parser("ofx")(path)

    assert len(output.transactions) == 2
    first, second = output.transactions
    assert first.txn_date == datetime.date(2024, 3, 2)
    assert first.direction == DEBIT
    assert first.debit_paise == 4210
    assert first.narration_raw == "WHOLE FOODS CARD 4411"
    assert first.currency == "USD"
    assert second.txn_date == datetime.date(2024, 3, 15)
    assert second.credit_paise == 150000
    assert output.statement_meta["closing_balance"] == 3457.90
    assert output.statement_meta["currency"] == "USD"


def test_qfx_uses_the_ofx_parser(tmp_path):
    path = _write(tmp_path / "statement.qfx", OFX1_SGML)
    assert get_parser("qfx") is get_parser("ofx")
    output = get_parser("qfx")(path)
    assert [t.debit_paise for t in output.transactions] == [125075, None]


def test_ofx_datetime_formats():
    """The OFX timestamp grammar, including the parts the normalizer cannot read."""
    assert ofx.parse_ofx_datetime("20240401") == datetime.date(2024, 4, 1)
    assert ofx.parse_ofx_datetime("20240401120000") == datetime.date(2024, 4, 1)
    assert ofx.parse_ofx_datetime("20240401120000.000") == datetime.date(2024, 4, 1)
    assert ofx.parse_ofx_datetime("20240401120000.000[-5:EST]") == datetime.date(2024, 4, 1)
    # The zone is not applied: a midnight posting stays on its own date rather
    # than sliding into the previous month.
    assert ofx.parse_ofx_datetime("20240401000000[+5:IST]") == datetime.date(2024, 4, 1)
    assert ofx.parse_ofx_datetime("20241332") is None
    assert ofx.parse_ofx_datetime("") is None


def test_ofx_without_an_ofx_element_is_rejected(tmp_path):
    path = _write(tmp_path / "notofx.ofx", "just some text\nwith no markup\n")
    with pytest.raises(ApiError) as excinfo:
        get_parser("ofx")(path)
    assert excinfo.value.code == "FILE_CORRUPT"


# ===========================================================================
# CAMT.053
# ===========================================================================

def test_camt053_parses_exact_values(tmp_path):
    path = _write(tmp_path / "camt.xml", CAMT053_XML)
    output = get_parser("xml_camt")(path)

    assert len(output.transactions) == 2
    first, second = output.transactions

    assert first.txn_date == datetime.date(2024, 4, 5)
    assert first.value_date == datetime.date(2024, 4, 6)
    assert first.direction == DEBIT
    assert first.debit_paise == 25075
    assert first.credit_paise is None
    assert first.balance_paise is None
    assert first.reference_number == "NTRY-1"
    assert first.narration_raw == "CARD PURCHASE MEDIAMARKT BERLIN"
    assert first.transaction_method == "PMNT-CCRD-POSD"
    assert first.currency == "EUR"

    assert second.txn_date == datetime.date(2024, 4, 10)
    assert second.direction == CREDIT
    assert second.credit_paise == 100100
    assert second.reference_number == "ASR-0002"
    assert second.narration_raw == "INVOICE 2024-114"

    meta = output.statement_meta
    assert meta["opening_balance"] == 1000.00
    assert meta["opening_balance_paise"] == 100000
    assert meta["closing_balance"] == 1750.25
    assert meta["closing_balance_paise"] == 175025
    assert meta["opening_balance_basis"] == "reported"
    assert meta["period_start"] == "2024-04-01"
    assert meta["period_end"] == "2024-04-30"
    assert meta["period_source"] == "declared"
    assert meta["account_number_masked"] == "******************3000"
    assert meta["currency"] == "EUR"

    assert output.continuity_passed is None
    assert output.rows_checked_for_continuity == 0
    assert W_CONTINUITY_UNVERIFIABLE in output.warning_codes


def test_camt053_older_schema_version_and_multi_currency(tmp_path):
    """camt.053.001.02 must work as well as .08 — matching is by local name."""
    path = _write(tmp_path / "fx.xml", CAMT053_MULTICCY)
    output = get_parser("xml_camt")(path)

    assert len(output.transactions) == 2
    assert [t.currency for t in output.transactions] == ["EUR", "USD"]
    assert output.transactions[0].direction == DEBIT
    assert output.transactions[0].debit_paise == 10000
    assert output.transactions[1].direction == CREDIT
    assert output.transactions[1].credit_paise == 25000
    assert W_MULTI_CURRENCY in output.warning_codes
    assert output.statement_meta["currencies"] == ["EUR", "USD"]
    assert output.statement_meta["account_number_masked"].endswith("6789")


def test_camt053_priced_txdtls_expand_one_row_each(tmp_path):
    path = _write(tmp_path / "batch.xml", CAMT053_BATCH)
    output = get_parser("xml_camt")(path)

    assert len(output.transactions) == 2, (
        "an Ntry whose TxDtls are individually priced is expanded")
    assert [t.debit_paise for t in output.transactions] == [12000, 18000]
    assert sum(t.amount_paise for t in output.transactions) == 30000, (
        "the expansion must still reconcile to the booked entry amount")
    assert [t.narration_raw for t in output.transactions] == ["SUPPLIER A", "SUPPLIER B"]


def test_camt_xxe_payload_is_not_resolved(tmp_path):
    """Customer-supplied XML must not be able to read the server's filesystem."""
    secret = tmp_path / "secret.txt"
    secret.write_text("TOPSECRET-ACCOUNT-KEY-98765")

    payload = f"""<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE Document [
  <!ENTITY xxe SYSTEM "file://{secret}">
]>
<Document xmlns="urn:iso:std:iso:20022:tech:xsd:camt.053.001.08">
  <BkToCstmrStmt><Stmt><Id>X</Id>
    <Ntry>
      <Amt Ccy="EUR">1.00</Amt>
      <CdtDbtInd>DBIT</CdtDbtInd>
      <BookgDt><Dt>2024-04-01</Dt></BookgDt>
      <AddtlNtryInf>&xxe;</AddtlNtryInf>
    </Ntry>
  </Stmt></BkToCstmrStmt>
</Document>
"""
    path = _write(tmp_path / "xxe.xml", payload)

    with pytest.raises(ApiError) as excinfo:
        get_parser("xml_camt")(path)

    error = excinfo.value
    assert error.code == "FILE_CORRUPT"
    assert "TOPSECRET" not in error.message
    assert "DTD" in error.message or "entities" in error.message


def test_camt_billion_laughs_is_refused(tmp_path):
    payload = """<?xml version="1.0"?>
<!DOCTYPE lolz [
 <!ENTITY lol "lol">
 <!ENTITY lol2 "&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;">
 <!ENTITY lol3 "&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;">
]>
<Document xmlns="urn:iso:std:iso:20022:tech:xsd:camt.053.001.08">
  <BkToCstmrStmt><Stmt><Id>&lol3;</Id></Stmt></BkToCstmrStmt>
</Document>
"""
    path = _write(tmp_path / "lol.xml", payload)
    with pytest.raises(ApiError) as excinfo:
        get_parser("xml_camt")(path)
    assert excinfo.value.code == "FILE_CORRUPT"


def test_camt_malformed_xml_is_file_corrupt(tmp_path):
    path = _write(tmp_path / "bad.xml",
                  '<?xml version="1.0"?><Document><BkToCstmrStmt><Stmt>')
    with pytest.raises(ApiError) as excinfo:
        get_parser("xml_camt")(path)
    assert excinfo.value.code == "FILE_CORRUPT"


def test_camt_valid_xml_that_is_not_camt_is_rejected(tmp_path):
    path = _write(tmp_path / "invoices.xml", "<invoices><invoice id='1'/></invoices>")
    with pytest.raises(ApiError) as excinfo:
        get_parser("xml_camt")(path)
    assert excinfo.value.code == "FILE_CORRUPT"
    assert "CAMT.053" in excinfo.value.message


# ===========================================================================
# Corrupt / empty / unsupported
# ===========================================================================

def test_empty_file_is_rejected_by_every_parser(tmp_path):
    for name, format_id in (("e.csv", "csv"), ("e.tsv", "tsv"), ("e.json", "json"),
                            ("e.ofx", "ofx"), ("e.xml", "xml_camt")):
        path = tmp_path / name
        path.write_bytes(b"")
        with pytest.raises(ApiError) as excinfo:
            get_parser(format_id)(str(path))
        assert excinfo.value.code == "FILE_EMPTY", f"{format_id} on an empty file"
        assert excinfo.value.status_code == 422


def test_corrupt_binary_named_csv_fails_with_a_code(tmp_path):
    path = tmp_path / "statement.csv"
    path.write_bytes(bytes(range(256)) * 8)
    with pytest.raises(ApiError) as excinfo:
        get_parser("csv")(str(path))
    # Either "no header row" or "unreadable"; what matters is that it is an
    # ApiError with a client-actionable code, not a library traceback.
    assert excinfo.value.code in ("PARSE_FAILED", "FILE_CORRUPT", "NO_TRANSACTIONS_FOUND")
    assert excinfo.value.status_code == 422


def test_csv_with_headers_but_no_rows_reports_no_transactions(tmp_path):
    path = _write(tmp_path / "headers_only.csv", ",".join(HEADER) + "\n")
    with pytest.raises(ApiError) as excinfo:
        get_parser("csv")(str(path))
    assert excinfo.value.code == "NO_TRANSACTIONS_FOUND"


def test_get_parser_rejects_unsupported_formats():
    with pytest.raises(ApiError) as excinfo:
        get_parser("zip")
    assert excinfo.value.code == "UNSUPPORTED_FILE_FORMAT"
    assert excinfo.value.status_code == 415
    assert "archive" in excinfo.value.message.lower()

    with pytest.raises(ApiError) as excinfo:
        get_parser("docx")
    assert excinfo.value.code == "UNSUPPORTED_FILE_FORMAT"

    with pytest.raises(ApiError) as excinfo:
        get_parser("wingdings")
    assert "supported formats are" in excinfo.value.message


def test_registry_is_complete_and_honest():
    ids = [entry["id"] for entry in SUPPORTED_FORMATS]
    assert ids == ["pdf", "xlsx", "xls", "csv", "tsv", "txt", "json",
                   "ofx", "qfx", "xml_camt"]
    for entry in SUPPORTED_FORMATS:
        assert entry["extensions"] and entry["mime_types"] and entry["description"]
        assert entry["notes"], f"{entry['id']} publishes no limitations"
        # Every published id must actually be dispatchable.
        assert callable(get_parser(entry["id"]))

    # The limitations that matter most are stated, not implied.
    def notes(format_id):
        for entry in SUPPORTED_FORMATS:
            if entry["id"] == format_id:
                return " ".join(entry["notes"])
        raise AssertionError(f"{format_id} is not published")

    assert "OCR" in notes("pdf")
    assert "xlrd" in notes("xls")
    assert "no per-transaction running balance" in notes("ofx")
    assert "camt.052" in notes("xml_camt")
    assert {e["id"] for e in UNSUPPORTED_FORMATS} >= {"zip", "xml_other"}


# ===========================================================================
# ingest.py
# ===========================================================================

def test_save_upload_streams_and_hashes(tmp_path):
    import hashlib

    data = make_delimited(tmp_path / "src.csv", ",")
    raw = open(data, "rb").read()
    upload = FakeUpload("april statement.csv", raw)

    ingested = save_upload(upload, max_bytes=1024 * 1024, tmp_dir=str(tmp_path))
    try:
        assert ingested.size_bytes == len(raw)
        assert ingested.sha256 == hashlib.sha256(raw).hexdigest()
        assert ingested.head_bytes == raw[:8192]
        assert os.path.isfile(ingested.path)
        # The stored name is generated, never the client's.
        assert "april" not in os.path.basename(ingested.path)
        assert ingested.detected.format == "csv"
        # The parser must work on what was written.
        assert len(get_parser("csv")(ingested.path).transactions) == 3
    finally:
        cleanup(ingested)


def test_save_upload_rejects_oversized_before_finishing(tmp_path):
    upload = FakeUpload("big.csv", b"x" * (3 * 1024 * 1024))
    with pytest.raises(ApiError) as excinfo:
        save_upload(upload, max_bytes=1024 * 1024, tmp_dir=str(tmp_path))

    assert excinfo.value.code == "FILE_TOO_LARGE"
    assert excinfo.value.status_code == 413
    assert excinfo.value.detail["max_bytes"] == 1024 * 1024
    # The partial write is not left behind.
    assert os.listdir(tmp_path) == []


def test_save_upload_rejects_empty_file(tmp_path):
    with pytest.raises(ApiError) as excinfo:
        save_upload(FakeUpload("empty.csv", b""), max_bytes=1024,
                    tmp_dir=str(tmp_path))
    assert excinfo.value.code == "FILE_EMPTY"
    assert os.listdir(tmp_path) == []


def test_save_upload_rejects_zip_archive(tmp_path):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("april.csv", ",".join(HEADER) + "\n")
    upload = FakeUpload("statements.zip", buffer.getvalue())

    with pytest.raises(ApiError) as excinfo:
        save_upload(upload, max_bytes=1024 * 1024, tmp_dir=str(tmp_path))

    assert excinfo.value.code == "UNSUPPORTED_FILE_FORMAT"
    assert excinfo.value.status_code == 415
    assert "archive" in excinfo.value.message.lower()
    assert os.listdir(tmp_path) == [], "the rejected archive is not left on disk"


def test_save_upload_accepts_xlsx_which_is_also_a_zip(tmp_path):
    """The zip rejection must not swallow OOXML workbooks."""
    path = make_xlsx(tmp_path / "book.xlsx")
    upload = FakeUpload("book.xlsx", open(path, "rb").read())
    ingested = save_upload(upload, max_bytes=1024 * 1024, tmp_dir=str(tmp_path))
    try:
        assert ingested.detected.format == "xlsx"
    finally:
        cleanup(ingested)


def test_save_upload_reports_extension_mismatch(tmp_path):
    upload = FakeUpload("statement.csv", b"%PDF-1.5\n%\xe2\xe3\xcf\xd3\ntrailer\n")
    ingested = save_upload(upload, max_bytes=1024 * 1024, tmp_dir=str(tmp_path))
    try:
        assert ingested.detected.format == "pdf"
        assert ingested.detected.mismatch is True
        assert ingested.detected.extension_format == "csv"
    finally:
        cleanup(ingested)


def test_filename_sanitisation_strips_traversal():
    assert sanitise_filename("../../etc/passwd") == "passwd"
    assert sanitise_filename("..\\..\\windows\\system32\\config") == "config"
    assert sanitise_filename("/absolute/path/statement.csv") == "statement.csv"
    assert sanitise_filename("") == "upload"
    assert sanitise_filename(None) == "upload"
    assert sanitise_filename("..") == "upload"
    assert "/" not in sanitise_filename("a/b/c.csv")
    long_name = "x" * 400 + ".csv"
    assert len(sanitise_filename(long_name)) <= 180


def test_traversal_filename_cannot_escape_the_temp_dir(tmp_path):
    upload = FakeUpload("../../../../tmp/evil.csv", b"Date,Description\n")
    ingested = save_upload(upload, max_bytes=1024, tmp_dir=str(tmp_path))
    try:
        assert os.path.realpath(ingested.path).startswith(os.path.realpath(str(tmp_path)))
        assert ingested.filename == "evil.csv"
    finally:
        cleanup(ingested)


def test_cleanup_is_idempotent(tmp_path):
    upload = FakeUpload("s.csv", b"Date,Description\n01/04/2024,X\n")
    ingested = save_upload(upload, max_bytes=1024, tmp_dir=str(tmp_path))
    directory = ingested.tmp_dir
    assert os.path.isdir(directory)

    cleanup(ingested)
    assert not os.path.exists(directory)
    cleanup(ingested)          # must not raise on a second call
    cleanup(None)              # nor on nothing at all


def test_ingested_file_exposes_extension(tmp_path):
    ingested = IngestedFile(path="/tmp/x/abc.csv", filename="Statement.CSV",
                            size_bytes=1, sha256="", head_bytes=b"")
    assert ingested.extension == ".csv"
