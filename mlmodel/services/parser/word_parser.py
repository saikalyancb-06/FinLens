import io
import re
import logging
import pandas as pd

logger = logging.getLogger(__name__)

def parse_word_file(file_content: bytes, filename: str = "") -> pd.DataFrame:
    """
    Word Document (.docx) Parser:
    Stage 1: Extract structured tables.
    Stage 2: Fallback to paragraph text extraction with regex transaction matching.
    """
    logger.info(f"Starting Word document parsing for {filename}")
    try:
        import docx
        doc = docx.Document(io.BytesIO(file_content))
        
        # Stage 1: Table Extraction
        table_rows = []
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    table_rows.append(cells)

        if table_rows:
            # Convert extracted table cells to DataFrame
            max_cols = max(len(r) for r in table_rows)
            padded = [r + [''] * (max_cols - len(r)) for r in table_rows]
            df_raw = pd.DataFrame(padded)
            # Use first row as header if it contains date/narration keywords
            header_str = " ".join([str(c).lower() for c in df_raw.iloc[0].values])
            if any(k in header_str for k in ['date', 'narration', 'particular', 'withdrawal', 'deposit', 'balance']):
                df_raw.columns = df_raw.iloc[0]
                df_raw = df_raw[1:]
            return df_raw

        # Stage 2: Paragraph text fallback
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        text_full = "\n".join(paragraphs)

        from services.parser.pdf_parser import _parse_raw_text_lines
        return _parse_raw_text_lines(text_full)

    except Exception as e:
        logger.error(f"Word document parsing error for {filename}: {e}")
        return pd.DataFrame()
